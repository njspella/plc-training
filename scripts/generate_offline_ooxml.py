#!/usr/bin/env python3
"""
Offline backup: PowerPoint + Word from scripts/training_data.json.

PowerPoint: prefers python-pptx if installed (layout aligned with web slide deck).
            Falls back to stdlib-only OOXML (valid DrawingML) so .pptx always generates.

Word: OOXML with Vermeer styling (navy #1e3a5f, gold #e8a820, Calibri) — stdlib only.
"""
from __future__ import annotations

import hashlib
import io
import time
import urllib.error
import urllib.request
from typing import Optional, Tuple
import json
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))
from commons_images import fetch_commons_bytes, url_for_images_src

JSON_PATH = SCRIPTS / "training_data.json"
EXPORTS = ROOT / "exports"
PPTX_DIR = EXPORTS / "powerpoint"
WORD_DIR = EXPORTS / "word"
# Cached downloads for Wikimedia / remote images (PowerPoint embeds local files)
IMAGE_CACHE_DIR = EXPORTS / ".offline_image_cache"

# Match css/styles.css :root
NAVY = (0x1E, 0x3A, 0x5F)
GOLD = (0xE8, 0xA8, 0x20)
# Web slide h2 uses --color-primary-dark (not navy)
PRIMARY_DARK = (0x0F, 0x24, 0x40)
BG_PAGE = "F8FAFC"
TEXT_MAIN = "1E293B"
TEXT_MUTED = "64748B"
WHITE = "FFFFFF"

KNOWLEDGE_ANSWERS = [
    "The PLC power supply converts incoming AC to low-voltage DC for the backplane and I/O. Field I/O is typically supplied 24 VDC (per cabinet design).",
    "N/O: open when de-energized, closes when true. N/C: closed when de-energized, opens when true.",
    "Likely causes: wiring fault (open, loose terminal, fuse) or NPN/PNP/common mismatch so the input circuit does not complete.",
    "OSSD = Output Signal Switching Device. Two independent channels so one fault cannot silently defeat the safety function.",
    "Same subnet as the PLC (e.g. 192.168.5.x with mask 255.255.255.0 for a 192.168.5.20 PLC).",
    "Communications → Who Active / Browse Chassis. Go online shortcut often Ctrl+D (confirm version).",
    "DN is ON when accumulated time reaches/exceeds preset — at 5000/5000 ms, DN is ON.",
    "Global tags are controller-wide; program-scoped tags are limited to their program unless exposed.",
    "Typical HMI levels: Operator, Maintenance/Technician, Administrator/Engineer (names vary).",
    "Level 2 — escalate servo/STO issues to Lead Technician / Noah Staudacher per policy.",
]

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


def safe_title_filename(title: str, max_len: int = 55) -> str:
    t = re.sub(r"[^\w\s-]+", "", str(title or "module"), flags=re.UNICODE)
    t = re.sub(r"[\s-]+", "_", t).strip("_") or "module"
    return t[:max_len]


def strip_html(s: str) -> str:
    if not s:
        return ""
    t = re.sub(r"<[^>]+>", "", str(s))
    t = t.replace("&nbsp;", " ")
    return " ".join(t.split()).strip()


def sanitize_ooxml_text(s: str, *, max_len: int | None = None) -> str:
    """
    Strip characters that break OOXML / XML 1.0 text (PowerPoint may refuse to open
    the file or show "repair" if these appear). Removes control chars and lone
    UTF-16 surrogates (invalid in XML character data).
    """
    if s is None:
        s = ""
    out: list[str] = []
    for ch in str(s):
        o = ord(ch)
        if o < 0x20 and o not in (0x9, 0xA, 0xD):
            continue
        if 0xD800 <= o <= 0xDFFF:
            continue
        if o in (0xFFFE, 0xFFFF):
            continue
        out.append(ch)
    t = "".join(out).strip()
    if not t:
        t = " "
    if max_len is not None and len(t) > max_len:
        t = t[:max_len]
    return t


def blocks_to_lines(blocks: list) -> list[str]:
    lines: list[str] = []

    def walk(blks):
        for b in blks or []:
            t = b.get("type")
            if t == "paragraph":
                lines.append(strip_html(b.get("text", "")))
            elif t == "heading":
                lines.append("→ " + strip_html(b.get("text", "")))
            elif t in ("list", "numbered", "steps"):
                for i, item in enumerate(b.get("items") or []):
                    prefix = f"{i + 1}. " if t == "numbered" else "• "
                    lines.append(prefix + strip_html(item))
            elif t == "callout":
                lines.append(f"[{b.get('variant', 'note').upper()}] {b.get('title', '')}: {strip_html(b.get('text', ''))}")
            elif t == "image":
                lines.append(f"[Figure] {strip_html(b.get('caption') or b.get('alt') or '')}")
            elif t == "table":
                hdrs = b.get("headers") or []
                lines.append("Table: " + " | ".join(hdrs))
                for row in b.get("rows") or []:
                    lines.append("  • " + " | ".join(str(c) for c in row))
            elif t == "sideBySide":
                walk(b.get("left") or [])
                walk(b.get("right") or [])

    walk(blocks)
    return [x for x in lines if x]


# Same keys as js/app.js SLIDE_IMAGE_FALLBACK — used when Commons files are not in images/ yet.
SLIDE_IMAGE_FALLBACK: dict[str, str] = {
    "images/cabinet_plc_cpu.jpg": "images/ps_cpu_module.png",
    "images/cabinet_power_supply_24v.jpg": "images/ps_plc_power_supply.png",
    "images/cabinet_industrial_ethernet.jpg": "images/ps_ethernet_board.png",
    "images/io_inductive_proximity.jpg": "images/inputs_sensors.png",
    "images/io_limit_switch_roller.jpg": "images/inputs_buttons.png",
    "images/io_operator_panel_pushbuttons.JPG": "images/inputs_buttons.png",
    "images/io_emergency_stop.jpg": "images/inputs_safety.png",
    "images/io_photoelectric.jpg": "images/inputs_sensors.png",
    "images/safety_laser_scanner_3d.jpg": "images/area_scanner.png",
    "images/safety_light_curtain.jpg": "images/inputs_safety.png",
    "images/safety_relay_module.jpg": "images/outputs_relay_light.png",
    "images/io_interposing_relay.jpg": "images/outputs_relay_light.png",
    "images/io_stack_light.jpg": "images/outputs_relay_light.png",
    "images/io_servo_motor.jpg": "images/outputs_motors.png",
    "images/io_servo_drive.jpg": "images/outputs_controller.png",
    "images/io_stepper_motor.jpg": "images/outputs_motors.png",
    "images/io_vfd.jpg": "images/outputs_motors.png",
    "images/solenoid_valve_coil.jpg": "images/outputs_relay_light.png",
    "images/rotary_encoder.jpg": "images/inputs_sensors.png",
    "images/dol_motor_starter.jpg": "images/outputs_motors.png",
    "images/hmi_screen.jpg": "images/image6.jpeg",
}


def resolve_image_src(src: str) -> Optional[Path]:
    """Return a local path to the image (project file or downloaded cache)."""
    src = (src or "").strip()
    if not src:
        return None
    if src.startswith(("http://", "https://")):
        IMAGE_CACHE_DIR.mkdir(parents=True, exist_ok=True)
        key = hashlib.sha256(src.encode("utf-8")).hexdigest()[:24]
        tail = src.split("?", 1)[0].split("/")[-1]
        ext = Path(tail).suffix.lower() if "." in tail else ""
        if ext not in (".jpg", ".jpeg", ".png", ".gif", ".webp"):
            ext = ".jpg"
        dest = IMAGE_CACHE_DIR / f"{key}{ext}"
        if dest.is_file():
            return dest
        try:
            req = urllib.request.Request(
                src,
                headers={"User-Agent": "plc-training-offline-export/1.0 (educational; +https://example.invalid)"},
            )
            with urllib.request.urlopen(req, timeout=60) as resp:
                data = resp.read()
            if not data:
                return None
            dest.write_bytes(data)
            return dest
        except (urllib.error.URLError, OSError, TimeoutError) as e:
            print(f"Warning: could not fetch image ({src[:72]}…): {e}", file=sys.stderr)
            return None
    p = Path(src)
    if not p.is_absolute():
        p = ROOT / p
    if p.is_file():
        return p
    # Slide assets under images/: pull from Commons when folder is empty (embeds real photos in .pptx)
    commons_url = url_for_images_src(src)
    if commons_url:
        data = fetch_commons_bytes(commons_url)
        if data:
            try:
                p.parent.mkdir(parents=True, exist_ok=True)
                p.write_bytes(data)
                return p
            except OSError:
                IMAGE_CACHE_DIR.mkdir(parents=True, exist_ok=True)
                cache_dest = IMAGE_CACHE_DIR / p.name
                try:
                    cache_dest.write_bytes(data)
                    return cache_dest
                except OSError:
                    pass
        else:
            print(
                f"Warning: could not download Commons image for {src} (offline or blocked)",
                file=sys.stderr,
            )
    alt = SLIDE_IMAGE_FALLBACK.get(src.strip())
    if alt:
        p2 = ROOT / alt
        if p2.is_file():
            return p2
    print(f"Warning: image not found: {src}", file=sys.stderr)
    return None


def build_lesson_content_slides_like_web(lesson: dict) -> list[dict]:
    """Mirror js buildSlides() for lesson.content: new slide on heading; split at 5 blocks (unless image)."""
    slides: list[dict] = []
    current: Optional[dict] = None
    for block in lesson.get("content") or []:
        bt = block.get("type")
        if bt in ("heading", "subheading"):
            if current and current.get("blocks"):
                slides.append(current)
            current = {"heading": block.get("text") or "", "blocks": []}
        else:
            if current is None:
                current = {"heading": "", "blocks": []}
            current["blocks"].append(block)
            if len(current["blocks"]) >= 5 and bt != "image":
                slides.append(current)
                current = {"heading": "", "blocks": []}
    if current and current.get("blocks"):
        slides.append(current)
    return slides


def pop_first_image_block(blocks: list) -> tuple[Optional[dict], list]:
    """
    Depth-first: remove the first image block (including inside sideBySide).
    Returns (image_block_or_none, remaining_blocks). Matches web slides that show a figure
    even when it is not blocks[0] (e.g. paragraph then image, or sideBySide with image).
    """
    if not blocks:
        return None, []
    b0 = blocks[0]
    t = b0.get("type")
    if t == "image":
        return b0, blocks[1:]
    if t == "sideBySide":
        left = list(b0.get("left") or [])
        right = list(b0.get("right") or [])
        img, nl = pop_first_image_block(left)
        if img is not None:
            new_side = {**b0, "left": nl, "right": right}
            return img, [new_side] + blocks[1:]
        img2, nr = pop_first_image_block(right)
        if img2 is not None:
            new_side = {**b0, "left": left, "right": nr}
            return img2, [new_side] + blocks[1:]
    img3, rest = pop_first_image_block(blocks[1:])
    if img3 is not None:
        return img3, [b0] + rest
    return None, blocks


def lesson_content_slide_specs(lesson: dict, lesson_title: str) -> list[tuple[str, list[str], Optional[Path]]]:
    """One export slide per web content slide; embed first figure on that slide (same slide as the website)."""
    out: list[tuple[str, list[str], Optional[Path]]] = []
    for slide in build_lesson_content_slides_like_web(lesson):
        heading = (slide.get("heading") or "").strip()
        title = heading if heading else lesson_title
        blocks = slide.get("blocks") or []
        if not blocks:
            continue
        img_block, rest_blocks = pop_first_image_block(blocks)
        if img_block is not None:
            img_path = resolve_image_src(img_block.get("src") or "")
            bullets = blocks_to_lines(rest_blocks)
            if not bullets:
                cap = strip_html(img_block.get("caption") or img_block.get("alt") or "")
                bullets = [cap] if cap else [" "]
        else:
            img_path = None
            bullets = blocks_to_lines(blocks)
        if not bullets:
            bullets = [" "]
        out.append((title, bullets, img_path))
    return out


# --- Word: styled OOXML (Vermeer) ---
def w_run(text: str, *, bold=False, size_half_pts=22, color=None) -> str:
    rpr = f'<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="{size_half_pts}"/>'
    if bold:
        rpr += "<w:b/>"
    if color:
        rpr += f'<w:color w:val="{color}"/>'
    rpr += "</w:rPr>"
    return f"<w:r>{rpr}<w:t xml:space=\"preserve\">{escape(text)}</w:t></w:r>"


def w_para_runs(runs_xml: str, *, ppr_inner: str = "") -> str:
    ppr = f"<w:pPr>{ppr_inner}</w:pPr>" if ppr_inner else ""
    return f"<w:p>{ppr}{runs_xml}</w:p>"


def w_hero_title(text: str) -> str:
    """Navy banner, white text — like .top-nav / slide deck."""
    ppr = (
        f'<w:pPr><w:shd w:val="clear" w:fill="{NAVY[0]:02X}{NAVY[1]:02X}{NAVY[2]:02X}"/>'
        f'<w:spacing w:before="120" w:after="200"/><w:ind w:left="360" w:right="360"/></w:pPr>'
    )
    r = w_run(text, bold=True, size_half_pts=56, color=WHITE)
    return w_para_runs(r, ppr_inner=ppr.replace("<w:pPr>", "").replace("</w:pPr>", ""))


def w_hero_sub(text: str) -> str:
    ppr = (
        f'<w:pPr><w:shd w:val="clear" w:fill="{NAVY[0]:02X}{NAVY[1]:02X}{NAVY[2]:02X}"/>'
        f'<w:spacing w:after="240"/><w:ind w:left="360" w:right="360"/></w:pPr>'
    )
    inner = ppr.replace("<w:pPr>", "").replace("</w:pPr>", "")
    r = w_run(text, bold=False, size_half_pts=22, color="E2E8F0")
    return w_para_runs(r, ppr_inner=inner)


def w_gold_rule() -> str:
    """Thin gold line."""
    return (
        f'<w:p><w:pPr><w:pBdr><w:bottom w:val="single" w:sz="24" w:space="1" w:color="{GOLD[0]:02X}{GOLD[1]:02X}{GOLD[2]:02X}"/></w:pBdr>'
        f'<w:spacing w:after="200"/></w:pPr></w:p>'
    )


def w_h2(text: str) -> str:
    """Section title: navy text + gold bottom border."""
    ppr = (
        f'<w:pPr><w:pBdr><w:bottom w:val="single" w:sz="18" w:space="2" w:color="{GOLD[0]:02X}{GOLD[1]:02X}{GOLD[2]:02X}"/></w:pBdr>'
        f'<w:spacing w:before="280" w:after="120"/></w:pPr>'
    )
    inner = ppr.replace("<w:pPr>", "").replace("</w:pPr>", "")
    r = w_run(text, bold=True, size_half_pts=36, color=f"{NAVY[0]:02X}{NAVY[1]:02X}{NAVY[2]:02X}")
    return w_para_runs(r, ppr_inner=inner)


def w_h3(text: str) -> str:
    r = w_run(text, bold=True, size_half_pts=28, color=f"{NAVY[0]:02X}{NAVY[1]:02X}{NAVY[2]:02X}")
    return w_para_runs(
        r,
        ppr_inner=f'<w:spacing w:before="200" w:after="80"/>',
    )


def w_body(text: str) -> str:
    r = w_run(text, size_half_pts=22, color=TEXT_MAIN)
    return w_para_runs(r, ppr_inner='<w:spacing w:after="120"/>')


def w_muted(text: str) -> str:
    r = w_run(text, size_half_pts=20, color=TEXT_MUTED)
    return w_para_runs(r, ppr_inner='<w:spacing w:after="160"/>')


def w_callout(text: str) -> str:
    """Light blue box like .callout.info on web."""
    ppr = (
        '<w:pPr><w:shd w:val="clear" w:fill="EFF6FF"/><w:spacing w:before="80" w:after="160"/>'
        '<w:ind w:left="360" w:right="360"/></w:pPr>'
    )
    inner = ppr.replace("<w:pPr>", "").replace("</w:pPr>", "")
    r = w_run(text, size_half_pts=22, color=TEXT_MAIN)
    return w_para_runs(r, ppr_inner=inner)


def build_document_xml(body_inner: str) -> bytes:
    doc = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
  xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<w:body>
{body_inner}
<w:sectPr>
  <w:pgSz w:w="12240" w:h="15840"/>
  <w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/>
  <w:pgBorders>
    <w:top w:val="single" w:sz="48" w:space="0" w:color="{GOLD[0]:02X}{GOLD[1]:02X}{GOLD[2]:02X}"/>
  </w:pgBorders>
</w:sectPr>
</w:body></w:document>"""
    return doc.encode("utf-8")


def pack_docx(document_xml: bytes, out: Path):
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, **_zipfile_open_kwargs()) as z:
        ooxml_zip_writestr(
            z,
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>
</Types>""",
        )
        ooxml_zip_writestr(
            z,
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""",
        )
        ooxml_zip_writestr(z, "word/document.xml", document_xml)
        ooxml_zip_writestr(
            z,
            "word/_rels/document.xml.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
        )
        ooxml_zip_writestr(
            z,
            "word/settings.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>',
        )


# --- PowerPoint: python-pptx (web-matched deck + optional images) ---
def pick_blank_layout(prs):
    """Built-in template layout indices differ by python-pptx version; prefer name 'Blank'."""
    for layout in prs.slide_layouts:
        name = (getattr(layout, "name", None) or "").strip().lower()
        if name == "blank":
            return layout
    try:
        return prs.slide_layouts[6]
    except IndexError:
        return prs.slide_layouts[-1]


def pack_pptx_lib(slides: list[tuple[str, list[str], Optional[Path]]], out: Path) -> None:
    """Slides aligned with web deck: .slide-accent-bar gradient, white card, h2 primary-dark + gold rule."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    try:
        prs.core_properties.title = "PLC Tabletop Training"
        prs.core_properties.subject = "Offline export"
    except Exception:
        pass

    blank = pick_blank_layout(prs)

    navy = RGBColor(*NAVY)
    gold = RGBColor(*GOLD)
    primary_dark = RGBColor(*PRIMARY_DARK)
    card_white = RGBColor(255, 255, 255)
    text = RGBColor(30, 41, 59)
    text_muted = RGBColor(100, 116, 139)
    border = RGBColor(226, 232, 240)

    # Match .slide-accent-bar (4px) + .slide.content-slide .slide-inner padding rhythm (tighter than before)
    bar_h = Inches(4.0 / 72.0)
    gap_after_bar = Inches(7.0 / 72.0)
    title_top = bar_h + gap_after_bar
    title_h = Inches(0.76)
    rule_h = Inches(3.0 / 72.0)
    gap_after_rule = Inches(8.0 / 72.0)
    body_top = title_top + title_h + Inches(0.03) + rule_h + gap_after_rule
    footer_h = Inches(6.0 / 72.0)
    body_h = Inches(7.5) - body_top - footer_h - Inches(0.14)

    for title, bullets, img_path in slides:
        slide = prs.slides.add_slide(blank)
        # Full slide: white card (.slide background)
        page = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        page.fill.solid()
        page.fill.fore_color.rgb = card_white
        page.line.width = Pt(0)
        # Top gradient bar: linear-gradient(90deg, primary, accent) — .slide-accent-bar
        accent_bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, bar_h
        )
        accent_bar.line.width = Pt(0)
        try:
            accent_bar.fill.gradient()
            accent_bar.fill.gradient_angle = 0.0
            gs = accent_bar.fill.gradient_stops
            gs[0].color.rgb = navy
            gs[1].color.rgb = gold
        except Exception:
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = navy
        # Gold rule under heading (h2 border-bottom: 3px solid accent)
        rule = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(0.45),
            title_top + title_h + Inches(0.03),
            Inches(9.1),
            rule_h,
        )
        rule.fill.solid()
        rule.fill.fore_color.rgb = gold
        rule.line.width = Pt(0)
        # Title: .slide.content-slide h2 — 24px bold, primary-dark
        tbox = slide.shapes.add_textbox(Inches(0.45), title_top, Inches(9.1), title_h)
        footer_y = float(prs.slide_height) - float(footer_h)
        footer_strip = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            0,
            footer_y,
            prs.slide_width,
            footer_h,
        )
        footer_strip.fill.solid()
        footer_strip.fill.fore_color.rgb = RGBColor(241, 245, 249)
        footer_strip.line.width = Pt(0)
        tf = tbox.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        p0.text = sanitize_ooxml_text(title or " ", max_len=500)
        p0.font.size = Pt(24)
        p0.font.bold = True
        p0.font.color.rgb = primary_dark
        p0.font.name = "Calibri"

        blist = bullets if bullets else [" "]
        use_image = img_path is not None and img_path.is_file()

        if use_image:
            # ~48% text / ~50% media — favor larger figures (aligned with web .slide-media-col)
            left_w = Inches(3.95)
            pic_left = Inches(4.48)
            pic_w = Inches(5.02)
            bbox = slide.shapes.add_textbox(Inches(0.48), body_top, left_w, body_h)
            btf = bbox.text_frame
            btf.word_wrap = True
            btf.margin_left = Inches(0.05)
            btf.margin_right = Inches(0.05)
            for i, line in enumerate(blist):
                bp = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
                bp.text = sanitize_ooxml_text(str(line), max_len=8000)
                bp.font.size = Pt(12)
                bp.font.name = "Calibri"
                bp.font.color.rgb = text
                bp.line_spacing = 1.12
                bp.space_after = Pt(2)
            try:
                pic = slide.shapes.add_picture(
                    str(img_path),
                    pic_left,
                    body_top,
                    width=pic_w,
                )
                try:
                    pic.line.color.rgb = border
                    pic.line.width = Pt(1)
                except Exception:
                    pass
            except (OSError, ValueError) as e:
                print(f"Warning: embed picture failed ({img_path}): {e}", file=sys.stderr)
                bbox2 = slide.shapes.add_textbox(pic_left, body_top, pic_w, Inches(1))
                bbox2.text_frame.paragraphs[0].text = "[Image could not be embedded]"
                bbox2.text_frame.paragraphs[0].font.size = Pt(10)
                bbox2.text_frame.paragraphs[0].font.color.rgb = text_muted
        else:
            bbox = slide.shapes.add_textbox(Inches(0.52), body_top, Inches(8.96), body_h)
            btf = bbox.text_frame
            btf.word_wrap = True
            btf.margin_left = Inches(0.06)
            btf.margin_right = Inches(0.06)
            for i, line in enumerate(blist):
                bp = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
                bp.text = sanitize_ooxml_text(str(line), max_len=8000)
                bp.font.size = Pt(12)
                bp.font.name = "Calibri"
                bp.font.color.rgb = text
                bp.line_spacing = 1.12
                bp.space_after = Pt(3)

    out.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".pptx", dir=str(out.parent))
    os.close(fd)
    tmp_path = Path(tmp)
    try:
        prs.save(str(tmp_path))
        with zipfile.ZipFile(tmp_path, "r") as z:
            if z.testzip() is not None:
                raise RuntimeError("generated .pptx failed zip integrity check")
            names = z.namelist()
            if "ppt/presentation.xml" not in names:
                raise RuntimeError("generated .pptx missing ppt/presentation.xml")
        os.replace(tmp_path, out)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise


def pack_pptx(slides: list[tuple[str, list[str], Optional[Path]]], out: Path, doc_title: str = "") -> None:
    if HAS_PPTX:
        pack_pptx_lib(slides, out)
    else:
        pack_pptx_stdlib(slides, out)


# --- PowerPoint: stdlib OOXML (valid runs — opens in desktop PowerPoint without python-pptx) ---
def _pptx_a_t(text: str) -> str:
    """Escape text for <a:t> and preserve spaces."""
    t = escape(sanitize_ooxml_text(text))
    return f'<a:t xml:space="preserve">{t}</a:t>'


def _pptx_title_run(text: str) -> str:
    """Web .slide.content-slide h2: 24px bold, --color-primary-dark."""
    pd = f"{PRIMARY_DARK[0]:02X}{PRIMARY_DARK[1]:02X}{PRIMARY_DARK[2]:02X}"
    return (
        f'<a:r><a:rPr sz="2400" b="1" lang="en-US">'
        f'<a:solidFill><a:srgbClr val="{pd}"/></a:solidFill>'
        f"</a:rPr>{_pptx_a_t(sanitize_ooxml_text(text, max_len=500))}</a:r>"
    )


# Slide geometry in EMUs (914400 per inch; matches p:sldSz).
# Web: .slide-accent-bar 4px + gap, h2 area, 3px gold rule, then body (.slide-inner padding)
_SLIDE_CY = 6858000
_ACCENT_BAR_CY = 50800  # 4px
_TITLE_GAP = 91440  # ~0.1" below bar (tighter than before)
_TITLE_X = 457200
_TITLE_Y = _ACCENT_BAR_CY + _TITLE_GAP
_TITLE_CX = 8229600
_TITLE_CY = 900000
_UNDERLINE_Y = _TITLE_Y + _TITLE_CY
_UNDERLINE_CY = 38100  # 3px gold rule
_BODY_GAP = 91440
_BODY_X = 457200
_BODY_Y = _UNDERLINE_Y + _UNDERLINE_CY + _BODY_GAP
_BODY_CX = 8229600
_FOOTER_CY = 76200  # ~5px subtle strip + room for gold hairline in fill
_BOTTOM_PAD = 45720
_BODY_CY = _SLIDE_CY - _FOOTER_CY - _BOTTOM_PAD - _BODY_Y
# Two-column slide (text | image) — wider figure column (~49% of content width)
_BODY_NARROW_CX = 3600000
_PIC_GAP = 101600
_PIC_X = _BODY_X + _BODY_NARROW_CX + _PIC_GAP
_PIC_Y = _BODY_Y
_CONTENT_RIGHT = _TITLE_X + _TITLE_CX
_PIC_CX = _CONTENT_RIGHT - _PIC_X
_PIC_CY = 5400000
_FOOTER_Y = _SLIDE_CY - _FOOTER_CY


def _pptx_body_line_text(line: str) -> str:
    """Avoid double bullets: lstStyle adds •; strip same prefix from content lines."""
    t = sanitize_ooxml_text(str(line), max_len=8000)
    if t.lstrip().startswith("•"):
        t = t.lstrip()[1:].lstrip()
    return t


def _pptx_body_paragraphs(lines: list[str]) -> str:
    """Body lines with bullets; ~12pt (was 11pt) to match denser web slide body."""
    body = TEXT_MAIN
    parts: list[str] = []
    for line in lines if lines else [" "]:
        parts.append(
            f'<a:p><a:pPr lvl="0" algn="l"><a:lnSpc><a:spcPct val="112000"/></a:lnSpc></a:pPr>'
            f'<a:r><a:rPr sz="1200" lang="en-US">'
            f'<a:solidFill><a:srgbClr val="{body}"/></a:solidFill>'
            f"</a:rPr>{_pptx_a_t(_pptx_body_line_text(line))}</a:r></a:p>"
        )
    return "".join(parts)


def build_slide_part_xml(title: str, bullets: list[str], *, image_embed_id: Optional[str] = None) -> str:
    """Title + content slide; optional embedded image (rId2) for stdlib .pptx without python-pptx."""
    title_xml = _pptx_title_run(title or " ")
    body_xml = _pptx_body_paragraphs(bullets if bullets else [" "])
    navy = f"{NAVY[0]:02X}{NAVY[1]:02X}{NAVY[2]:02X}"
    gold = f"{GOLD[0]:02X}{GOLD[1]:02X}{GOLD[2]:02X}"
    accent_bar_sp = f"""<p:sp>
  <p:nvSpPr><p:cNvPr id="5" name="AccentBar"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="0" y="0"/><a:ext cx="9144000" cy="{_ACCENT_BAR_CY}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:gradFill rotWithShape="0">
      <a:gsLst>
        <a:gs pos="0"><a:srgbClr val="{navy}"/></a:gs>
        <a:gs pos="100000"><a:srgbClr val="{gold}"/></a:gs>
      </a:gsLst>
      <a:lin ang="0" scaled="0" flip="none"/>
    </a:gradFill>
  </p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
</p:sp>"""
    footer_bar_sp = f"""<p:sp>
  <p:nvSpPr><p:cNvPr id="7" name="FooterStrip"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="0" y="{_FOOTER_Y}"/><a:ext cx="9144000" cy="{_FOOTER_CY}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:solidFill><a:srgbClr val="F1F5F9"/></a:solidFill>
    <a:ln><a:noFill/></a:ln>
  </p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
</p:sp>"""
    underline_sp = f"""<p:sp>
  <p:nvSpPr><p:cNvPr id="6" name="HeadingUnderline"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{_TITLE_X}" y="{_UNDERLINE_Y}"/><a:ext cx="{_TITLE_CX}" cy="{_UNDERLINE_CY}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:solidFill><a:srgbClr val="{gold}"/></a:solidFill>
  </p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
</p:sp>"""
    lst_style = f"""<a:lstStyle>
  <a:lvl1pPr marL="457200" indent="-228600" algn="l">
    <a:buChar char="•"/>
    <a:defRPr sz="1200"><a:solidFill><a:srgbClr val="{TEXT_MAIN}"/></a:solidFill></a:defRPr>
  </a:lvl1pPr>
</a:lstStyle>"""
    title_sppr = f"""<p:spPr>
  <a:xfrm>
    <a:off x="{_TITLE_X}" y="{_TITLE_Y}"/>
    <a:ext cx="{_TITLE_CX}" cy="{_TITLE_CY}"/>
  </a:xfrm>
  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
</p:spPr>"""
    if image_embed_id:
        body_sppr = f"""<p:spPr>
  <a:xfrm>
    <a:off x="{_BODY_X}" y="{_BODY_Y}"/>
    <a:ext cx="{_BODY_NARROW_CX}" cy="{_BODY_CY}"/>
  </a:xfrm>
  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
</p:spPr>"""
        pic_xml = f"""<p:pic>
  <p:nvPicPr>
    <p:cNvPr id="4" name="Figure" descr=""/>
    <p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>
    <p:nvPr/>
  </p:nvPicPr>
  <p:blipFill>
    <a:blip r:embed="{image_embed_id}"/>
    <a:stretch><a:fillRect/></a:stretch>
  </p:blipFill>
  <p:spPr>
    <a:xfrm>
      <a:off x="{_PIC_X}" y="{_PIC_Y}"/>
      <a:ext cx="{_PIC_CX}" cy="{_PIC_CY}"/>
    </a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
  </p:spPr>
</p:pic>"""
        tail = f"""</p:txBody></p:sp>
{pic_xml}
</p:spTree></p:cSld></p:sld>"""
    else:
        body_sppr = f"""<p:spPr>
  <a:xfrm>
    <a:off x="{_BODY_X}" y="{_BODY_Y}"/>
    <a:ext cx="{_BODY_CX}" cy="{_BODY_CY}"/>
  </a:xfrm>
  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
</p:spPr>"""
        tail = """</p:txBody></p:sp>
</p:spTree></p:cSld></p:sld>"""
    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
  xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
  xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:cSld><p:spTree>
<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>
{accent_bar_sp}
{footer_bar_sp}
<p:sp><p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
{title_sppr}
<p:txBody><a:bodyPr wrap="square" anchor="t" vert="horz"/><a:lstStyle/><a:p>{title_xml}</a:p></p:txBody></p:sp>
{underline_sp}
<p:sp><p:nvSpPr><p:cNvPr id="3" name="Content"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph idx="1" type="body"/></p:nvPr></p:nvSpPr>
{body_sppr}
<p:txBody><a:bodyPr wrap="square" anchor="t" vert="horz" lIns="91440" tIns="45720" rIns="91440" bIns="45720"/>{lst_style}{body_xml}{tail}
"""


# Theme + full slide master — PowerPoint/Outlook reject packages without theme, notesSz, sldMasterIdLst, etc.
THEME1_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Office Theme">
  <a:themeElements>
    <a:clrScheme name="Office">
      <a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>
      <a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
      <a:dk2><a:srgbClr val="1F497D"/></a:dk2>
      <a:lt2><a:srgbClr val="EEECE1"/></a:lt2>
      <a:accent1><a:srgbClr val="4F81BD"/></a:accent1>
      <a:accent2><a:srgbClr val="C0504D"/></a:accent2>
      <a:accent3><a:srgbClr val="9BBB59"/></a:accent3>
      <a:accent4><a:srgbClr val="8064A2"/></a:accent4>
      <a:accent5><a:srgbClr val="4BACC6"/></a:accent5>
      <a:accent6><a:srgbClr val="F79646"/></a:accent6>
      <a:hlink><a:srgbClr val="0000FF"/></a:hlink>
      <a:folHlink><a:srgbClr val="800080"/></a:folHlink>
    </a:clrScheme>
    <a:fontScheme name="Office">
      <a:majorFont><a:latin typeface="Calibri"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont>
      <a:minorFont><a:latin typeface="Calibri"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont>
    </a:fontScheme>
    <a:fmtScheme name="Office">
      <a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst>
      <a:lnStyleLst><a:ln w="9525"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst>
      <a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst>
      <a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst>
    </a:fmtScheme>
  </a:themeElements>
  <a:objectDefaults/>
  <a:extraClrSchemeLst/>
</a:theme>
"""

SLIDE_MASTER_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sldMaster xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
 xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:cSld>
  <p:bg><p:bgPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></p:bgPr></p:bg>
  <p:spTree>
    <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
    <p:grpSpPr/>
  </p:spTree>
</p:cSld>
<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>
<p:sldLayoutIdLst>
  <p:sldLayoutId id="2147483649" r:id="rId1"/>
</p:sldLayoutIdLst>
<p:txStyles>
  <p:titleStyle>
    <a:lvl1pPr algn="ctr"><a:defRPr sz="4400" b="1"><a:latin typeface="Calibri"/></a:defRPr></a:lvl1pPr>
  </p:titleStyle>
  <p:bodyStyle>
    <a:lvl1pPr><a:defRPr sz="2800"><a:latin typeface="Calibri"/></a:defRPr></a:lvl1pPr>
  </p:bodyStyle>
</p:txStyles>
</p:sldMaster>
"""

# Title + body placeholder geometry must exist on the layout or Impress omits body in PDF export.
SLIDE_LAYOUT_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sldLayout xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
 xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" type="obj" preserve="1">
<p:cSld name="Title and Content"><p:spTree>
<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>
<p:sp>
  <p:nvSpPr><p:cNvPr id="2" name="Title Placeholder"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{_TITLE_X}" y="{_TITLE_Y}"/><a:ext cx="{_TITLE_CX}" cy="{_TITLE_CY}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
  </p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
</p:sp>
<p:sp>
  <p:nvSpPr><p:cNvPr id="3" name="Content Placeholder"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph idx="1" type="body"/></p:nvPr></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{_BODY_X}" y="{_BODY_Y}"/><a:ext cx="{_BODY_CX}" cy="{_BODY_CY}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
  </p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
</p:sp>
</p:spTree></p:cSld></p:sldLayout>
"""

CORE_PROPS_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
 xmlns:dc="http://purl.org/dc/elements/1.1/">
  <dc:title>PLC Tabletop Training</dc:title>
  <dc:creator>plc-training export</dc:creator>
</cp:coreProperties>
"""

def app_props_xml(slide_count: int) -> str:
    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">
  <Application>plc-training</Application>
  <Slides>{slide_count}</Slides>
</Properties>
"""


def _zipfile_open_kwargs():
    """ZIP deflate level for OOXML packages (avoid strict_timestamps=False → 1980 dates; some mail scanners flag that)."""
    return {"compresslevel": 6}


def ooxml_zip_writestr(z: zipfile.ZipFile, arcname: str, data: bytes | str) -> None:
    """Write one member with a current timestamp (Python defaults to 1980-01-01; Gmail/antivirus may treat that as suspicious)."""
    if isinstance(data, str):
        data = data.encode("utf-8")
    zi = zipfile.ZipInfo(arcname)
    zi.compress_type = zipfile.ZIP_DEFLATED
    zi.date_time = time.localtime()[:6]
    zi.create_system = 0
    zi.external_attr = 0
    z.writestr(zi, data, compresslevel=6)


def _content_type_for_media_part(arc: str) -> str:
    ext = Path(arc).suffix.lower()
    return {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".gif": "image/gif",
        ".webp": "image/webp",
    }.get(ext, "image/png")


def _collect_pptx_embedded_media(
    slides: list[tuple[str, list[str], Optional[Path]]],
) -> tuple[list[Optional[str]], dict[str, bytes]]:
    """Per-slide ppt/media/... path or None; arcname -> image bytes (deduped by file path)."""
    path_to_arc: dict[str, str] = {}
    arc_to_bytes: dict[str, bytes] = {}
    per_slide: list[Optional[str]] = []
    mi = 0
    for _t, _b, img in slides:
        if img is not None and img.is_file():
            key = str(img.resolve())
            if key in path_to_arc:
                per_slide.append(path_to_arc[key])
                continue
            ext = img.suffix.lower()
            if ext not in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
                ext = ".png"
            arc = f"ppt/media/image{mi}{ext}"
            mi += 1
            path_to_arc[key] = arc
            arc_to_bytes[arc] = img.read_bytes()
            per_slide.append(arc)
        else:
            per_slide.append(None)
    return per_slide, arc_to_bytes


def _slide_rels_xml(has_image: bool, media_basename: Optional[str]) -> str:
    base = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>"""
    if not has_image or not media_basename:
        return base + "\n</Relationships>"
    return (
        base
        + f'\n<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/{media_basename}"/>'
        + "\n</Relationships>"
    )


def pack_pptx_stdlib(slides: list[tuple[str, list[str], Optional[Path]]], out: Path) -> None:
    """Build a minimal .pptx with zipfile only; embeds local/cached images without python-pptx."""
    slide_rows: list[tuple[str, list[str], Optional[Path]]] = list(slides)
    n = len(slide_rows)
    if n == 0:
        slide_rows = [("Empty", ["No content"], None)]
        n = 1

    slide_media, media_bytes = _collect_pptx_embedded_media(slide_rows)

    # [Content_Types].xml
    ct_parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
        '<Default Extension="xml" ContentType="application/xml"/>',
        '<Default Extension="png" ContentType="image/png"/>',
        '<Default Extension="jpg" ContentType="image/jpeg"/>',
        '<Default Extension="jpeg" ContentType="image/jpeg"/>',
        '<Default Extension="gif" ContentType="image/gif"/>',
        '<Default Extension="webp" ContentType="image/webp"/>',
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>',
        '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>',
        '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>',
        '<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>',
        '<Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>',
        '<Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>',
    ]
    for i in range(1, n + 1):
        ct_parts.append(
            f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        )
    for arc in sorted(media_bytes.keys()):
        ct_parts.append(f'<Override PartName="/{arc}" ContentType="{_content_type_for_media_part(arc)}"/>')
    ct_parts.append("</Types>")
    content_types = "\n".join(ct_parts)

    # presentation.xml.rels: rId1 = master, rId2.. = slides
    rels_lines = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">',
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="slideMasters/slideMaster1.xml"/>',
    ]
    for i in range(n):
        rels_lines.append(
            f'<Relationship Id="rId{i + 2}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{i + 1}.xml"/>'
        )
    rels_lines.append("</Relationships>")
    pres_rels = "\n".join(rels_lines)

    # presentation.xml
    sld_ids = []
    for i in range(n):
        sld_ids.append(f'<p:sldId id="{256 + i}" r:id="rId{i + 2}"/>')
    # notesSz + sldSz + sldMasterIdLst are required for many Microsoft readers (incl. Outlook preview)
    presentation_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
  xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
  xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" saveSubsetFonts="1">
<p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>
<p:sldIdLst>{"".join(sld_ids)}</p:sldIdLst>
<p:sldSz cx="9144000" cy="6858000"/>
<p:notesSz cx="6858000" cy="9144000"/>
<p:defaultTextStyle/>
</p:presentation>
"""

    slide_master_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="../theme/theme1.xml"/>
</Relationships>"""

    slide_layout_to_master_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="../slideMasters/slideMaster1.xml"/>
</Relationships>"""

    root_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>"""

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, **_zipfile_open_kwargs()) as z:
        ooxml_zip_writestr(z, "[Content_Types].xml", content_types)
        ooxml_zip_writestr(z, "_rels/.rels", root_rels)
        ooxml_zip_writestr(z, "docProps/core.xml", CORE_PROPS_XML)
        ooxml_zip_writestr(z, "docProps/app.xml", app_props_xml(n))
        ooxml_zip_writestr(z, "ppt/presentation.xml", presentation_xml.encode("utf-8"))
        ooxml_zip_writestr(z, "ppt/_rels/presentation.xml.rels", pres_rels.encode("utf-8"))
        ooxml_zip_writestr(z, "ppt/slideMasters/slideMaster1.xml", SLIDE_MASTER_XML.encode("utf-8"))
        ooxml_zip_writestr(z, "ppt/slideMasters/_rels/slideMaster1.xml.rels", slide_master_rels.encode("utf-8"))
        ooxml_zip_writestr(z, "ppt/theme/theme1.xml", THEME1_XML.encode("utf-8"))
        ooxml_zip_writestr(z, "ppt/slideLayouts/slideLayout1.xml", SLIDE_LAYOUT_XML.encode("utf-8"))
        ooxml_zip_writestr(z, "ppt/slideLayouts/_rels/slideLayout1.xml.rels", slide_layout_to_master_rels.encode("utf-8"))
        for arc, raw in media_bytes.items():
            ooxml_zip_writestr(z, arc, raw)
        for idx, (title, bl, _) in enumerate(slide_rows, start=1):
            media_arc = slide_media[idx - 1]
            embed = "rId2" if media_arc else None
            ooxml_zip_writestr(
                z,
                f"ppt/slides/slide{idx}.xml",
                build_slide_part_xml(title, list(bl), image_embed_id=embed).encode("utf-8"),
            )
            base_name = Path(media_arc).name if media_arc else None
            ooxml_zip_writestr(
                z,
                f"ppt/slides/_rels/slide{idx}.xml.rels",
                _slide_rels_xml(bool(media_arc), base_name).encode("utf-8"),
            )

    out.parent.mkdir(parents=True, exist_ok=True)
    raw = buf.getvalue()
    fd, tmp = tempfile.mkstemp(suffix=".pptx", dir=str(out.parent))
    os.close(fd)
    tmp_path = Path(tmp)
    try:
        tmp_path.write_bytes(raw)
        with zipfile.ZipFile(tmp_path, "r") as z:
            if z.testzip() is not None:
                raise RuntimeError("stdlib .pptx failed zip integrity check")
            if "ppt/presentation.xml" not in z.namelist():
                raise RuntimeError("stdlib .pptx missing ppt/presentation.xml")
        os.replace(tmp_path, out)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise


def write_scenarios_docx(td: dict, path: Path):
    parts = [
        w_hero_title("Hands-On Scenarios"),
        w_hero_sub(f"{td.get('title', '')} · v{td.get('version', '')} · {td.get('revisionDate', '')}"),
        w_gold_rule(),
        w_muted("Bench exercises — instructor-staged faults. Timing and success criteria as listed."),
    ]
    for sc in td.get("scenarios") or []:
        parts.append(w_h2(f"Scenario {sc.get('id')}: {sc.get('title', '')}"))
        parts.append(w_body(f"Objective: {strip_html(sc.get('objective', ''))}"))
        parts.append(w_muted(f"Time limit: {sc.get('timeLimit', '')}"))
        parts.append(w_h3("Setup (instructor)"))
        parts.append(w_body(strip_html(sc.get("setup", ""))))
        parts.append(w_h3("Steps"))
        for s in sc.get("steps") or []:
            parts.append(w_body(f"• {s}"))
        parts.append(w_h3("Success criteria"))
        parts.append(w_callout(strip_html(sc.get("successCriteria", ""))))
        parts.append(w_body(""))
    pack_docx(build_document_xml("".join(parts)), path)


def write_escalation_docx(td: dict, path: Path):
    parts = [
        w_hero_title("Escalation Procedures"),
        w_hero_sub("PLC & automation support — match training site policy"),
        w_gold_rule(),
    ]
    esc = td.get("escalation") or {}
    for lev in esc.get("levels") or []:
        parts.append(w_h2(f"Level {lev.get('level')}: {lev.get('title', '')}"))
        parts.append(w_muted(f"Tier: {lev.get('color', '')}"))
        for item in lev.get("items") or []:
            parts.append(w_body(f"• {item}"))
    parts.append(w_h2("Document before escalating"))
    for d in esc.get("documentation") or []:
        parts.append(w_body(f"• {d}"))
    pack_docx(build_document_xml("".join(parts)), path)


def write_knowledge_trainee(td: dict, path: Path):
    qs = td.get("knowledgeCheck") or []
    parts = [
        w_hero_title("Written Knowledge Check"),
        w_hero_sub(f"Trainee copy · {len(qs)} questions · PLC Tabletop Training"),
        w_gold_rule(),
        w_muted("Name: _________________________   Date: ______________   Score: _____ / " + str(len(qs))),
        w_body("Answer in complete sentences unless your instructor directs otherwise."),
    ]
    for i, q in enumerate(qs, 1):
        parts.append(w_h2(f"Question {i}"))
        parts.append(w_body(strip_html(q.get("question", ""))))
        parts.append(w_body("_" * 72))
        parts.append(w_body("_" * 72))
        parts.append(w_body(""))
    pack_docx(build_document_xml("".join(parts)), path)


def write_knowledge_key(td: dict, path: Path):
    qs = td.get("knowledgeCheck") or []
    parts = [
        w_hero_title("Knowledge Check — Answer Key"),
        w_hero_sub("INSTRUCTOR ONLY — do not distribute to trainees"),
        w_gold_rule(),
        w_callout("Remove or separate this file before printing class copies."),
    ]
    for i, q in enumerate(qs, 1):
        ans = KNOWLEDGE_ANSWERS[i - 1] if i <= len(KNOWLEDGE_ANSWERS) else "(Draft from materials.)"
        parts.append(w_h2(f"Question {i}"))
        parts.append(w_body(strip_html(q.get("question", ""))))
        parts.append(w_h3("Suggested answer"))
        parts.append(w_body(ans))
        parts.append(w_body(""))
    pack_docx(build_document_xml("".join(parts)), path)


def module_to_slide_specs(td: dict, mod: dict) -> list[tuple[str, list[str], Optional[Path]]]:
    """Slides as (title, bullet lines, optional image path) — images align with web training."""
    slides: list[tuple[str, list[str], Optional[Path]]] = []
    mid = mod.get("id")
    slides.append(
        (
            f"Module {mid}: {mod.get('title', '')}",
            [
                f"{td.get('title', '')} · {mod.get('format', '')} · {mod.get('hours', '')} h",
                f"v{td.get('version', '')} ({td.get('revisionDate', '')})",
            ],
            None,
        )
    )
    if mod.get("objectives"):
        slides.append(("Learning objectives", [strip_html(x) for x in mod["objectives"]], None))
    d = strip_html(mod.get("description", ""))
    if d:
        slides.append(("Overview", [d], None))
    for lesson in mod.get("lessons") or []:
        lt = lesson.get("title", "")
        summ = strip_html(lesson.get("summary", ""))
        slides.append((f"Lesson: {lt}", [summ] if summ else [], None))
        slides.extend(lesson_content_slide_specs(lesson, lt))
    for demo in mod.get("demoBreaks") or []:
        slides.append(
            (
                f"Lab: {demo.get('title', '')}",
                [
                    f"After: {demo.get('after', '')}",
                    f"Duration: {demo.get('duration', '')}",
                    strip_html(demo.get("description", "")),
                    strip_html(demo.get("activity", "")),
                ],
                None,
            )
        )
    for q in mod.get("quiz") or []:
        opts = " | ".join(strip_html(o) for o in (q.get("options") or []))
        slides.append(("Quiz", [strip_html(q.get("question", "")), opts], None))
    return slides


def module_to_slides(td: dict, mod: dict) -> list[tuple[str, list[str]]]:
    """Text-only slides for PDF/HTML generators (figures become [Figure] lines)."""
    out: list[tuple[str, list[str]]] = []
    for title, bullets, img in module_to_slide_specs(td, mod):
        if img is not None:
            cap = " ".join(bullets).strip() if bullets else ""
            out.append((title, [f"[Figure] {cap}"] if cap else ["[Figure]"]))
        else:
            out.append((title, bullets))
    return out


EMAIL_ZIP_README = """PLC Training — offline exports (read me first)
=====================================

Email attachments (Gmail, Outlook, etc.)
----------------------------------------
• Do NOT open .pptx or .docx from the browser preview. Download first.
• Best: attach the single file PLC_Training_Offline_Exports.zip (this archive).
• After download, right-click the ZIP → Extract All (Windows) or double-click and drag files out (Mac).
• Open the extracted .pptx files with Microsoft PowerPoint (desktop), not the browser.
• The pdf/ folder holds PDF versions of each PowerPoint (from LibreOffice when installed), or text PDFs / printable HTML as fallback — use any PDF viewer; Gmail handles PDFs more reliably than raw .pptx.

If Gmail bundled several attachments into one ZIP, extract that outer ZIP first, then extract this one.

If files still fail: upload the ZIP to Google Drive and use Share → link (avoids mail re-encoding).

Generated by scripts/generate_offline_ooxml.py
"""


def _zip_date_time_from_path(path: Path) -> tuple:
    try:
        return time.localtime(path.stat().st_mtime)[:6]
    except OSError:
        return time.localtime()[:6]


def distribution_zip_writestr(
    z: zipfile.ZipFile, arcname: str, data: bytes, *, date_time: Optional[Tuple[int, ...]] = None
) -> None:
    zi = zipfile.ZipInfo(arcname)
    zi.compress_type = zipfile.ZIP_DEFLATED
    zi.date_time = date_time if date_time is not None else time.localtime()[:6]
    zi.create_system = 0
    zi.external_attr = 0
    z.writestr(zi, data, compresslevel=6)


def zip_offline_exports(exports_root: Path) -> Optional[Path]:
    """One ZIP of all .pptx, .docx, and optional .pdf for email; avoids per-file MIME quirks and Gmail 'download all' confusion."""
    pptx_dir = exports_root / "powerpoint"
    word_dir = exports_root / "word"
    pdf_dir = exports_root / "pdf"
    html_dir = exports_root / "html"
    out_zip = exports_root / "PLC_Training_Offline_Exports.zip"

    files: list[tuple[Path, str]] = []
    if pptx_dir.is_dir():
        for p in sorted(pptx_dir.glob("*.pptx")):
            files.append((p, f"powerpoint/{p.name}"))
    if word_dir.is_dir():
        for p in sorted(word_dir.glob("*.docx")):
            files.append((p, f"word/{p.name}"))
    if pdf_dir.is_dir():
        for p in sorted(pdf_dir.glob("*.pdf")):
            files.append((p, f"pdf/{p.name}"))
    if html_dir.is_dir():
        for p in sorted(html_dir.glob("*.html")):
            files.append((p, f"html/{p.name}"))
    if not files:
        return None

    out_zip.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_zip.with_suffix(out_zip.suffix + ".tmp")
    try:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, **_zipfile_open_kwargs()) as z:
            distribution_zip_writestr(z, "README-Email.txt", EMAIL_ZIP_README.encode("utf-8"))
            for path, arcname in files:
                distribution_zip_writestr(z, arcname, path.read_bytes(), date_time=_zip_date_time_from_path(path))
        os.replace(tmp, out_zip)
    except Exception:
        tmp.unlink(missing_ok=True)
        raise
    return out_zip


def main():
    if not JSON_PATH.exists():
        print("Run: node scripts/export_training_data.js", file=sys.stderr)
        sys.exit(1)
    td = json.loads(JSON_PATH.read_text(encoding="utf-8")).get("TrainingData")
    if not td:
        sys.exit(1)

    PPTX_DIR.mkdir(parents=True, exist_ok=True)
    WORD_DIR.mkdir(parents=True, exist_ok=True)

    if HAS_PPTX:
        print("PowerPoint: using python-pptx (Vermeer layout + images when URLs resolve).", file=sys.stderr)
    else:
        print(
            "PowerPoint: stdlib OOXML (web-style accent bar + heading; embeds photos when cached). Optional: pip install -r scripts/requirements-offline.txt for python-pptx (richer layout + image borders).",
            file=sys.stderr,
        )

    for mod in td.get("modules") or []:
        slides = module_to_slide_specs(td, mod)
        if not slides:
            slides = [("Empty", ["No content"], None)]
        fname = f"Module_{mod.get('id')}_{safe_title_filename(mod.get('title', 'module'))}.pptx"
        outp = PPTX_DIR / fname
        pack_pptx(slides, outp, mod.get("title", ""))
        print("Wrote", outp)

    ov: list[tuple[str, list[str], Optional[Path]]] = [
        (td.get("title", "PLC Training"), [td.get("subtitle", ""), td.get("revisionDate", "")], None)
    ]
    for mod in td.get("modules") or []:
        lines = [strip_html(mod.get("description", ""))]
        lines += [strip_html(x) for x in (mod.get("objectives") or [])]
        ov.append((f"Module {mod.get('id')}: {mod.get('title', '')}", [x for x in lines if x], None))
    out_ov = PPTX_DIR / "ALL_MODULES_Overview.pptx"
    pack_pptx(ov, out_ov, "Overview")
    print("Wrote", out_ov)

    write_scenarios_docx(td, WORD_DIR / "Hands_On_Scenarios.docx")
    print("Wrote", WORD_DIR / "Hands_On_Scenarios.docx")
    write_escalation_docx(td, WORD_DIR / "Escalation_Procedures.docx")
    print("Wrote", WORD_DIR / "Escalation_Procedures.docx")
    write_knowledge_trainee(td, WORD_DIR / "Knowledge_Check_Trainee.docx")
    print("Wrote", WORD_DIR / "Knowledge_Check_Trainee.docx")
    write_knowledge_key(td, WORD_DIR / "Knowledge_Check_Answer_Key.docx")
    print("Wrote", WORD_DIR / "Knowledge_Check_Answer_Key.docx")
    print("Done:", EXPORTS)
    # Distribution ZIP is built after PDF step by generate_offline_pdf.py (so pdf/ is included).


if __name__ == "__main__":
    main()
